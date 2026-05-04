"""
PPTX Extractor — 슬라이드 단위 + 표 분리 + 발표자 노트.

각 슬라이드 = PageUnit 1개 (본문)
각 표      = PageUnit 1개 (parent_page_index = 본문 슬라이드 인덱스)
발표자 노트 = 본문 PageUnit의 raw_metadata["notes"]에 포함

PART 헤더(예: "PART 03 우리는 어떻게 다른가") 패턴 인식 → section_path 추적.
"""
from __future__ import annotations

import re
from pathlib import Path

from pptx import Presentation

from app.knowledge.chunking.page_unit import PageUnit


# PART 헤더 패턴 (예: "PART 01   우리는 왜 바뀌었나")
_PART_RE = re.compile(r"^(PART\s*\d+)\b", re.MULTILINE)


def _table_to_markdown(rows: list[list[str]]) -> str:
    if not rows:
        return ""
    col_count = max(len(r) for r in rows)
    if col_count == 0:
        return ""

    def _norm(cells: list[str]) -> list[str]:
        padded = (cells + [""] * col_count)[:col_count]
        return [c.replace("|", "\\|").replace("\n", " ").strip() for c in padded]

    header = _norm(rows[0])
    lines = ["| " + " | ".join(header) + " |"]
    lines.append("|" + "|".join(["---"] * col_count) + "|")
    for row in rows[1:]:
        lines.append("| " + " | ".join(_norm(row)) + " |")
    return "\n".join(lines)


def _extract_text_from_shape(shape) -> str:
    if shape.has_text_frame:
        return (shape.text_frame.text or "").strip()
    if hasattr(shape, "text") and shape.text:
        return shape.text.strip()
    return ""


def extract_pptx(path: Path) -> list[PageUnit]:
    """PPTX → list[PageUnit].

    각 슬라이드:
      1. 본문 PageUnit 생성 (제목 + body shapes 텍스트 + 노트)
      2. 슬라이드 안의 모든 표는 별도 PageUnit으로 분리 (parent_page_index 부착)

    section_path: PART N 헤더가 슬라이드 본문에 나오면 갱신, 다음 슬라이드까지 유지.
    """
    prs = Presentation(str(path))
    units: list[PageUnit] = []
    current_section = ""  # PART 헤더 추적

    for slide_idx, slide in enumerate(prs.slides, start=1):
        # ─── 1. 제목 추출 ────────────────────────────────
        title = ""
        title_shape = None
        try:
            title_shape = slide.shapes.title
        except Exception:
            title_shape = None
        if title_shape is not None:
            title = _extract_text_from_shape(title_shape)

        # ─── 2. 본문 텍스트 + 표 분리 ─────────────────────
        body_parts: list[str] = []
        tables_in_slide: list[list[list[str]]] = []

        for shape in slide.shapes:
            if title_shape is not None and shape is title_shape:
                continue
            # 표 처리
            if getattr(shape, "has_table", False):
                rows = []
                for row in shape.table.rows:
                    cells = [(cell.text or "").strip() for cell in row.cells]
                    if any(cells):
                        rows.append(cells)
                if rows:
                    tables_in_slide.append(rows)
                continue
            # 본문 텍스트
            text = _extract_text_from_shape(shape)
            if text:
                body_parts.append(text)

        body_text = "\n\n".join(body_parts).strip()

        # ─── 3. 발표자 노트 ───────────────────────────────
        notes_text = ""
        try:
            if slide.has_notes_slide:
                notes_text = (slide.notes_slide.notes_text_frame.text or "").strip()
        except Exception:
            pass

        # ─── 4. PART 헤더 갱신 (본문에서 패턴 검색) ─────
        full_searchable = f"{title}\n{body_text}"
        m = _PART_RE.search(full_searchable)
        if m:
            current_section = m.group(1)

        # ─── 5. 본문 PageUnit 생성 ────────────────────────
        # 본문이 비어있고 제목만 있어도 PageUnit 생성 (표지·섹션 구분 슬라이드도 추적 가치 있음)
        main_text_parts = []
        if title:
            main_text_parts.append(f"# {title}")
        if body_text:
            main_text_parts.append(body_text)
        main_text = "\n\n".join(main_text_parts) if main_text_parts else ""

        if main_text or tables_in_slide:
            units.append(PageUnit(
                unit_index=slide_idx,
                unit_type="slide",
                title=title,
                section_path=current_section,
                text=main_text,
                is_table=False,
                table_index=None,
                parent_page_index=None,
                raw_metadata={
                    "notes": notes_text,
                    "shape_count": len(slide.shapes),
                },
            ))

        # ─── 6. 각 표 별 PageUnit ────────────────────────
        for t_idx, rows in enumerate(tables_in_slide, start=1):
            table_md = _table_to_markdown(rows)
            units.append(PageUnit(
                unit_index=slide_idx * 1000 + t_idx,  # 본문과 다른 unit_index
                unit_type="table",
                title=f"{title} — 표 {t_idx}" if title else f"슬라이드 {slide_idx} 표 {t_idx}",
                section_path=current_section,
                text=table_md,
                is_table=True,
                table_index=t_idx,
                parent_page_index=slide_idx,  # 본문 슬라이드와 연결
                raw_metadata={"row_count": len(rows)},
            ))

    return units
